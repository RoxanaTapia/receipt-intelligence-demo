"""Spreadsheet download of the receipts listed under Your receipts."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from datetime import date
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import main
from app.examples import ExampleReceipt
from app.receipt_export import guard_text, line_rows, render_csv, render_xlsx
from app.sample_receipt import ITEMS
from app.session import IMPORTS_COOKIE

SEED_DIR = Path(__file__).resolve().parents[1] / "seed"
HEADER = [
    "date",
    "merchant",
    "item",
    "quantity",
    "category",
    "amount",
    "receipt_total",
    "file_name",
]
_FORMULA_TRIGGERS = ("=SUM(A1)", "+1", "-1", "@cmd", "\t=1", "\r=1")


def _seed_line_count() -> int:
    manifest = json.loads((SEED_DIR / "examples.json").read_text(encoding="utf-8"))
    total = 0
    for entry in manifest:
        receipt = json.loads((SEED_DIR / entry["file"]).read_text(encoding="utf-8"))
        total += len(receipt["line_items"])
    return total


def _csv_rows(content: bytes) -> list[list[str]]:
    assert content.startswith(b"\xef\xbb\xbf")
    text = content.decode("utf-8-sig")
    assert "\n" not in text.replace("\r\n", "")
    return list(csv.reader(io.StringIO(text)))


def _write_live(data_dir: Path, name: str, receipt: dict) -> None:
    (data_dir / name).write_text(json.dumps(receipt), encoding="utf-8")


def _receipt(
    day: str,
    merchant: str,
    file_name: str,
    items: list[str],
    *,
    price: float = 1.0,
) -> ExampleReceipt:
    return ExampleReceipt(
        id=file_name,
        label=merchant,
        file_name=file_name,
        receipt={
            "date": day,
            "merchant": merchant,
            "total": price * len(items),
            "line_items": [
                {
                    "description": item,
                    "quantity": 1,
                    "price": price,
                    "category": "food",
                }
                for item in items
            ],
        },
    )


def test_csv_matches_seed_receipts_and_starts_with_bom(client) -> None:
    response = client.get("/export/receipts.csv")
    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="receipts.csv"'
    rows = _csv_rows(response.content)
    text = response.content.decode("utf-8-sig")
    assert rows[0] == HEADER
    assert len(rows) == 1 + _seed_line_count()
    assert "$" not in text
    assert "DC" not in text
    assert rows[1] == [
        "2026-07-12",
        "Lumen Pantry",
        "Oat milk",
        "1",
        "drinks",
        "2.40",
        "14.20",
        "2026-07-12_lumen-pantry_aaaaaaaa.json",
    ]
    pasta = next(row for row in rows if row[2] == "Pasta shells")
    assert pasta[3] == "2"
    assert pasta[5] == "3.20"
    assert pasta[6] == "21.40"
    harbor = [row[2] for row in rows if row[1] == "Harbor Home"]
    assert harbor == [
        "Kitchen sponge set",
        "Glass cleaner",
        "Pasta shells",
        "Tomato sauce",
        "Sparkling water",
    ]


def test_csv_umlaut_round_trips(client, data_dir: Path) -> None:
    name = "2026-05-01_kase-shop_aaaa1111.json"
    _write_live(
        data_dir,
        name,
        {
            "date": "2026-05-01",
            "merchant": "Käse Shop",
            "total": 3.50,
            "line_items": [
                {
                    "description": "Käse",
                    "quantity": 1,
                    "price": 3.50,
                    "category": "food",
                }
            ],
        },
    )
    client.cookies.set(IMPORTS_COOKIE, name)
    content = client.get("/export/receipts.csv").content
    assert "Käse".encode("utf-8") in content
    rows = _csv_rows(content)
    assert [
        "2026-05-01",
        "Käse Shop",
        "Käse",
        "1",
        "food",
        "3.50",
        "3.50",
        name,
    ] in rows


def test_csv_quotes_commas_and_keeps_crlf(client, data_dir: Path) -> None:
    name = "2026-05-02_quoted-stall_bbbb2222.json"
    _write_live(
        data_dir,
        name,
        {
            "date": "2026-05-02",
            "merchant": 'North, "Main" Stall',
            "total": 2.00,
            "line_items": [
                {
                    "description": "Oat milk, large",
                    "quantity": 1,
                    "price": 2.00,
                    "category": "drinks",
                }
            ],
        },
    )
    client.cookies.set(IMPORTS_COOKIE, name)
    content = client.get("/export/receipts.csv").content
    assert b'""Main""' in content
    rows = _csv_rows(content)
    quoted = next(row for row in rows if row[7] == name)
    assert quoted[1] == 'North, "Main" Stall'
    assert quoted[2] == "Oat milk, large"


def test_csv_prefixes_formula_text_and_keeps_negative_amount(
    client, data_dir: Path
) -> None:
    name = "2026-04-01_guard-shop_cccc3333.json"
    _write_live(
        data_dir,
        name,
        {
            "date": "2026-04-01",
            "merchant": "Guard Shop",
            "total": -0.50,
            "line_items": [
                {
                    "description": "=SUM(A1)",
                    "quantity": 1,
                    "price": -0.50,
                    "category": "food",
                }
            ],
        },
    )
    client.cookies.set(IMPORTS_COOKIE, name)
    content = client.get("/export/receipts.csv").content
    text = content.decode("utf-8-sig")
    assert "'=SUM(A1)" in text
    assert ",=SUM(A1)," not in text
    row = next(item for item in _csv_rows(content) if item[1] == "Guard Shop")
    assert row[2] == "'=SUM(A1)"
    assert row[5] == "-0.50"
    assert row[6] == "-0.50"
    assert not row[5].startswith("'")


def test_xlsx_sheet_matches_seed_receipts(client) -> None:
    response = client.get("/export/receipts.xlsx")
    assert response.status_code == 200
    assert (
        response.headers["content-disposition"] == 'attachment; filename="receipts.xlsx"'
    )
    sheet = load_workbook(io.BytesIO(response.content))["Line items"]
    assert [cell.value for cell in sheet[1]] == HEADER
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:H{sheet.max_row}"
    assert sheet.max_row == 1 + _seed_line_count()
    assert isinstance(sheet["A2"].value, date)
    assert sheet["A2"].value.year == 2026
    assert sheet["A2"].value.month == 7
    assert sheet["A2"].value.day == 12
    assert sheet["D2"].value == 1
    assert isinstance(sheet["D2"].value, (int, float))
    assert sheet["F2"].value == 2.4
    assert isinstance(sheet["F2"].value, (int, float))
    assert sheet["F2"].number_format == "0.00"
    assert sheet["G2"].value == 14.2
    assert sheet["G2"].number_format == "0.00"
    for data_row in sheet.iter_rows(min_row=2):
        assert isinstance(data_row[3].value, (int, float))
        assert isinstance(data_row[5].value, (int, float))
        assert isinstance(data_row[6].value, (int, float))
        assert data_row[5].number_format == "0.00"
        assert data_row[6].number_format == "0.00"
        assert data_row[5].data_type != "f"
    for cell in _cells(sheet):
        assert cell.data_type != "f"


def test_xlsx_cells_are_never_formulas(client, data_dir: Path) -> None:
    name = "2026-04-02_formula-guard_dddd4444.json"
    _write_live(
        data_dir,
        name,
        {
            "date": "2026-04-02",
            "merchant": "Guard Shop",
            "total": -0.50,
            "line_items": [
                {
                    "description": trigger,
                    "quantity": 1,
                    "price": -0.50,
                    "category": "food",
                }
                for trigger in _FORMULA_TRIGGERS
            ],
        },
    )
    client.cookies.set(IMPORTS_COOKIE, name)
    csv_rows = _csv_rows(client.get("/export/receipts.csv").content)
    items = [row[2] for row in csv_rows]
    for trigger in _FORMULA_TRIGGERS:
        assert f"'{trigger}" in items

    content = client.get("/export/receipts.xlsx").content
    # The workbook stores the apostrophe in front of the carriage return.
    # openpyxl turns that CR into a line feed when it loads the cell.
    assert b"'\r=1" in _xlsx_xml(content)
    sheet = load_workbook(io.BytesIO(content))["Line items"]
    seen: list[object] = []
    for cell in _cells(sheet):
        assert cell.data_type != "f"
        if isinstance(cell.value, str):
            assert cell.value[:1] not in "=+-@\t\r"
        seen.append(cell.value)
    for trigger in _FORMULA_TRIGGERS:
        expected = "'\n=1" if trigger == "\r=1" else f"'{trigger}"
        assert expected in seen
    amounts = [
        row[5].value
        for row in sheet.iter_rows(min_row=2)
        if row[2].value == "'=SUM(A1)"
    ]
    assert amounts == [-0.5]
    assert all(
        row[5].data_type == "n" and row[5].number_format == "0.00"
        for row in sheet.iter_rows(min_row=2)
        if isinstance(row[2].value, str) and row[2].value.startswith("'")
    )


def test_export_skips_another_sessions_live_import(
    client, fake_n8n, data_dir: Path
) -> None:
    stranger = "2026-09-04_stranger-mart_00112233.json"
    _write_live(
        data_dir,
        stranger,
        {
            "date": "2026-09-04",
            "merchant": "Stranger Mart",
            "total": 4.0,
            "line_items": [
                {
                    "description": "Hidden cola",
                    "quantity": 1,
                    "price": 4.0,
                    "category": "drinks",
                }
            ],
        },
    )
    client.post("/ingest-sample", data={"start": "2026-07-01", "end": "2026-08-31"})
    rows = _csv_rows(client.get("/export/receipts.csv").content)
    items = [row[2] for row in rows]
    merchants = [row[1] for row in rows]
    assert "Bananas" in items
    assert fake_n8n.filename in [row[7] for row in rows]
    assert "Hidden cola" not in items
    assert "Stranger Mart" not in merchants
    assert len(rows) == 1 + _seed_line_count() + len(ITEMS)

    values = _xlsx_values(client.get("/export/receipts.xlsx").content)
    assert "Bananas" in values
    assert "Hidden cola" not in values
    assert "Stranger Mart" not in values

    other = TestClient(main.app)
    other_rows = _csv_rows(other.get("/export/receipts.csv").content)
    assert "Bananas" not in [row[2] for row in other_rows]
    assert "Hidden cola" not in [row[2] for row in other_rows]
    assert len(other_rows) == 1 + _seed_line_count()
    other_values = _xlsx_values(other.get("/export/receipts.xlsx").content)
    assert "Bananas" not in other_values
    assert "Hidden cola" not in other_values


def test_page_renders_download_buttons(client) -> None:
    html = client.get("/").text
    block = html.split('id="your-receipts"', 1)[1].split("</section>", 1)[0]
    assert 'class="btn" href="export/receipts.xlsx"' in block
    assert 'class="btn-secondary" href="export/receipts.csv"' in block
    assert block.index("Download Excel") < block.index("Download CSV (for Sheets)")


def test_download_buttons_follow_public_url_prefix(client, monkeypatch) -> None:
    monkeypatch.setattr(main, "ROOT_PATH", "/app")
    html = client.get("/").text
    block = html.split('id="your-receipts"', 1)[1]
    assert 'class="btn" href="/app/export/receipts.xlsx"' in block
    assert 'class="btn-secondary" href="/app/export/receipts.csv"' in block


def test_download_buttons_hidden_without_receipts(client, monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setattr(main, "SEED_DIR", tmp_path / "no-seeds")
    html = client.get("/").text
    assert "No receipts loaded yet." in html
    assert "Download Excel" not in html
    assert "Download CSV" not in html
    assert _csv_rows(client.get("/export/receipts.csv").content) == [HEADER]


def test_export_does_not_call_the_receipt_api(client, monkeypatch) -> None:
    def boom(*_args, **_kwargs):
        raise AssertionError("receipt API must not be called")

    monkeypatch.setattr(main.api, "health", boom)
    monkeypatch.setattr(main.api, "summary", boom)
    monkeypatch.setattr(main.api, "ask", boom)
    assert client.get("/export/receipts.csv").status_code == 200
    assert client.get("/export/receipts.xlsx").status_code == 200


def test_rows_sort_by_date_merchant_file_and_line_order() -> None:
    rows = line_rows(
        [
            _receipt("2026-08-02", "Mango", "m.json", ["Third"]),
            _receipt("2026-08-02", "Alpha", "b.json", ["Second"]),
            _receipt("2026-08-01", "Zucchini", "z.json", ["Beta", "Alpha"]),
            _receipt("2026-08-02", "Alpha", "a.json", ["First"]),
        ],
        [],
    )
    assert [(row.date_text, row.merchant, row.file_name, row.item) for row in rows] == [
        ("2026-08-01", "Zucchini", "z.json", "Beta"),
        ("2026-08-01", "Zucchini", "z.json", "Alpha"),
        ("2026-08-02", "Alpha", "a.json", "First"),
        ("2026-08-02", "Alpha", "b.json", "Second"),
        ("2026-08-02", "Mango", "m.json", "Third"),
    ]


def test_guard_text_prefixes_only_formula_like_text() -> None:
    for raw in _FORMULA_TRIGGERS:
        assert guard_text(raw) == f"'{raw}"
    assert guard_text("Oat milk") == "Oat milk"
    assert guard_text("") == ""
    risky = line_rows(
        [_receipt("2026-01-01", "Shop", "a.json", ["=SUM(A1)"], price=-0.5)],
        [],
    )
    row = _csv_rows(render_csv(risky))[1]
    assert row[2] == "'=SUM(A1)"
    assert row[5] == "-0.50"
    sheet = load_workbook(io.BytesIO(render_xlsx(risky)))["Line items"]
    assert sheet["C2"].value == "'=SUM(A1)"
    assert sheet["C2"].data_type != "f"
    assert sheet["F2"].value == -0.5
    assert sheet["F2"].data_type == "n"
    for cell in _cells(sheet):
        assert cell.data_type != "f"


def _cells(sheet):
    for row in sheet.iter_rows():
        yield from row


def _xlsx_values(content: bytes) -> list[object]:
    sheet = load_workbook(io.BytesIO(content)).active
    return [cell.value for cell in _cells(sheet)]


def _xlsx_xml(content: bytes) -> bytes:
    blob = bytearray()
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        for name in package.namelist():
            if name.endswith(".xml"):
                blob.extend(package.read(name))
    return bytes(blob)
