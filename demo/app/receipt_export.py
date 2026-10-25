"""CSV and Excel download of the receipts shown under Your receipts.

The rows come from this browser's session catalog — seeds plus this
cookie's live imports — so the file matches the list on the page.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

COLUMNS = (
    "date",
    "merchant",
    "item",
    "quantity",
    "category",
    "amount",
    "receipt_total",
    "file_name",
)
SHEET_NAME = "Line items"
AMOUNT_FORMAT = "0.00"

# A leading character from this set makes Excel and Sheets treat the cell as a formula.
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


@dataclass(frozen=True)
class LineRow:
    """One exported line item. Text fields are already formula-safe."""

    date: date | None
    date_text: str
    merchant: str
    item: str
    quantity: int | float
    category: str
    amount: float
    receipt_total: float
    file_name: str


def guard_text(value: str) -> str:
    """Prefix formula-like text so a spreadsheet stores it as text.

    Numeric cells must not pass through here: a negative amount has to stay
    a number, not a quoted string.
    """
    if value.startswith(_FORMULA_STARTS):
        return f"'{value}"
    return value


def line_rows(examples: list[Any], live_imports: list[Any]) -> list[LineRow]:
    """One spreadsheet row per line item.

    ``amount`` is the line's ``price`` (already the line total), rounded to
    two decimal places. Rows are ordered by date, merchant, file name, then
    the original line order.
    """
    ranked: list[tuple[tuple[str, str, str, int], LineRow]] = []
    for item in [*examples, *live_imports]:
        ranked.extend(_rows_for_receipt(item))
    ranked.sort(key=lambda pair: pair[0])
    return [row for _, row in ranked]


def render_csv(rows: list[LineRow]) -> bytes:
    """UTF-8 CSV with a BOM, comma separator, and CRLF line endings."""
    buffer = io.StringIO(newline="")
    writer = csv.writer(
        buffer,
        delimiter=",",
        quotechar='"',
        doublequote=True,
        lineterminator="\r\n",
        quoting=csv.QUOTE_MINIMAL,
    )
    writer.writerow(COLUMNS)
    for row in rows:
        writer.writerow(
            [
                row.date_text,
                row.merchant,
                row.item,
                row.quantity,
                row.category,
                _amount_text(row.amount),
                _amount_text(row.receipt_total),
                row.file_name,
            ]
        )
    return buffer.getvalue().encode("utf-8-sig")


def render_xlsx(rows: list[LineRow]) -> bytes:
    """Excel workbook: one sheet, frozen bold header, amounts as numbers."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    _write_header(sheet)
    for index, row in enumerate(rows, start=2):
        _write_row(sheet, index, row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{sheet.max_row}"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _rows_for_receipt(
    item: Any,
) -> list[tuple[tuple[str, str, str, int], LineRow]]:
    receipt = item.receipt if isinstance(item.receipt, dict) else {}
    file_name = str(getattr(item, "file_name", "") or "")
    merchant = _plain(receipt.get("merchant"))
    raw_date = _plain(receipt.get("date"))
    parsed = _parse_date(raw_date)
    lines = _lines(receipt)
    total = _receipt_total(receipt, lines)
    date_key = parsed.isoformat() if parsed else raw_date
    date_text = date_key if parsed else guard_text(raw_date)
    rows: list[tuple[tuple[str, str, str, int], LineRow]] = []
    for index, line in enumerate(lines):
        row = LineRow(
            date=parsed,
            date_text=date_text,
            merchant=guard_text(merchant),
            item=guard_text(_plain(line.get("description"))),
            quantity=_quantity(line.get("quantity")),
            category=guard_text(_plain(line.get("category"))),
            amount=_cents(line.get("price")),
            receipt_total=total,
            file_name=guard_text(file_name),
        )
        rows.append(((date_key, merchant, file_name, index), row))
    return rows


def _write_header(sheet: Any) -> None:
    bold = Font(bold=True)
    for column, name in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, column, name)
        cell.font = bold


def _write_row(sheet: Any, index: int, row: LineRow) -> None:
    if row.date is not None:
        sheet.cell(index, 1, row.date)
    else:
        sheet.cell(index, 1, row.date_text)
    sheet.cell(index, 2, row.merchant)
    sheet.cell(index, 3, row.item)
    sheet.cell(index, 4, row.quantity)
    sheet.cell(index, 5, row.category)
    amount = sheet.cell(index, 6, row.amount)
    amount.number_format = AMOUNT_FORMAT
    total = sheet.cell(index, 7, row.receipt_total)
    total.number_format = AMOUNT_FORMAT
    sheet.cell(index, 8, row.file_name)


def _receipt_total(receipt: dict[str, Any], lines: list[dict[str, Any]]) -> float:
    raw = receipt.get("total")
    if raw is not None:
        try:
            return round(float(raw), 2)
        except (TypeError, ValueError):
            pass
    total = 0.0
    for line in lines:
        total = round(total + _cents(line.get("price")), 2)
    return total


def _cents(value: Any) -> float:
    """Two decimal places, as a number — no currency symbol."""
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return 0.0


def _quantity(value: Any) -> int | float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return 0
    if number.is_integer():
        return int(number)
    return number


def _amount_text(value: float) -> str:
    return f"{value:.2f}"


def _lines(receipt: dict[str, Any]) -> list[dict[str, Any]]:
    raw = receipt.get("line_items") or []
    if not isinstance(raw, list):
        return []
    return [line for line in raw if isinstance(line, dict)]


def _plain(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _parse_date(value: str) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        return None
